import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader

from doc_harness.exporters import write_docx, write_pdf
from doc_harness.sources import IndexedSource


SAMPLE = "# Weekly Reading\n\n## Key Ideas\n\nUseful source [1].\n\n## Sources\n\n[1] Example — https://example.com\n"


class ExporterTests(unittest.TestCase):
    def test_docx_pdf_contain_readable_content(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            word = folder / "reading.docx"
            pdf = folder / "reading.pdf"
            write_docx(SAMPLE, word)
            write_pdf(SAMPLE, pdf)
            self.assertIn("Useful source [1].", "\n".join(p.text for p in Document(word).paragraphs))
            self.assertIn("Useful source [1].", PdfReader(pdf).pages[0].extract_text())

    def test_word_navigation_contains_only_collection_sections_and_articles(self):
        source = IndexedSource(
            "https://example.org/blog/a-real-article",
            "https://example.org/blog/a-real-article",
            "# A Real Article\n\n## An actual section\n\nA complete paragraph.\nAnother line.\n\n"
            "## Share\n\nA fallback source may include this text.",
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "reading.docx"
            write_docx("unused", path, [source], "An edition summary [1].")
            doc = Document(path)
            headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
            self.assertEqual(headings, [
                "Edition Summary", "Contents", "Complete Articles",
                "Source 1: A Real Article", "Original Sources",
            ])
            self.assertTrue(any(p.style.name == "Article Subheading" for p in doc.paragraphs))
            self.assertIn("A complete paragraph. Another line.", [p.text for p in doc.paragraphs])
            self.assertIn('w:anchor="source-1"', doc._element.xml)
            with ZipFile(path) as archive:
                settings = archive.read("word/settings.xml").decode("utf-8")
                self.assertIn('w:name="compatibilityMode"', settings)
                self.assertIn('w:val="15"', settings)
                self.assertIn('w:percent="100"', settings)
