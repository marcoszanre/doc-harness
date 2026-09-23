"""Three exporters from one approved Markdown draft, with bounded local writes."""

from __future__ import annotations

import html
import hashlib
import os
import re
import tempfile
from datetime import date
from pathlib import Path
from urllib.parse import urlsplit

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

from .sources import IndexedSource, read_file
from .workspace import Workspace, atomic_text
from .workflow import compose, load_snapshots, review, source_title


def lines(markdown: str):
    for line in markdown.splitlines():
        text = line.strip()
        if not text:
            continue
        level = len(text) - len(text.lstrip("#"))
        if level in (1, 2, 3) and text[level:level + 1] == " ":
            yield ("heading", level, text[level:].strip())
        elif text.startswith(("- ", "* ")):
            yield ("bullet", 0, text[2:])
        else:
            yield ("paragraph", 0, text)


def plain(text: str) -> str:
    return re.sub(r"(?<!\w)[*`_]+|[*`_]+(?!\w)", "", text)


def word_blocks(markdown: str):
    paragraph: list[str] = []
    for line in markdown.splitlines():
        text = line.strip()
        heading = re.match(r"^(#{1,3})\s+(.+)$", text)
        if not text or heading or text.startswith(("- ", "* ")):
            if paragraph:
                yield "paragraph", " ".join(paragraph)
                paragraph = []
            if heading:
                yield "heading", heading.group(2)
            elif text.startswith(("- ", "* ")):
                yield "bullet", text[2:]
        else:
            paragraph.append(text)
    if paragraph:
        yield "paragraph", " ".join(paragraph)


def _temp_output(path: Path, create) -> None:
    fd, name = tempfile.mkstemp(prefix=".doc-harness-", suffix=path.suffix, dir=path.parent)
    os.close(fd)
    temporary = Path(name)
    try:
        create(temporary)
        if not temporary.is_file() or temporary.stat().st_size == 0:
            raise ValueError("Exporter produced an empty file.")
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _word_link(paragraph, label: str, destination: str, external: bool = False) -> None:
    link = OxmlElement("w:hyperlink")
    if external:
        relationship = paragraph.part.relate_to(destination, RT.HYPERLINK, is_external=True)
        link.set(qn("r:id"), relationship)
    else:
        link.set(qn("w:anchor"), destination)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2155BC")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def _word_body(
    doc: Document, text: str, link_citations: bool = False, skip_title: str = "",
) -> None:
    for kind, content in word_blocks(text):
        content = plain(content)
        if kind == "heading":
            if content.casefold() == skip_title.casefold():
                continue
            paragraph = doc.add_paragraph(style="Article Subheading")
        elif kind == "bullet":
            paragraph = doc.add_paragraph(style="List Bullet")
        else:
            paragraph = doc.add_paragraph()
        if link_citations:
            for part in re.split(r"(\[\d+\])", content):
                if re.fullmatch(r"\[\d+\]", part):
                    _word_link(paragraph, part, f"source-{part[1:-1]}")
                elif part:
                    paragraph.add_run(part)
        else:
            paragraph.add_run(content)


def write_docx(
    markdown: str, path: Path, sources: list[IndexedSource] | None = None,
    summary: str = "",
) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(35, 43, 52)
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.space_after = Pt(7)
    for name, size in (("Title", 21), ("Heading 1", 16), ("Heading 2", 13)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(30, 46, 65)
        style.paragraph_format.space_before = Pt(16)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True
    subheading = doc.styles.add_style("Article Subheading", WD_STYLE_TYPE.PARAGRAPH)
    subheading.base_style = normal
    subheading.font.bold = True
    subheading.font.size = Pt(11.5)
    subheading.font.color.rgb = RGBColor(40, 58, 76)
    subheading.paragraph_format.space_before = Pt(12)
    subheading.paragraph_format.space_after = Pt(5)
    subheading.paragraph_format.keep_with_next = True
    compat = doc.settings.element.find(qn("w:compat"))
    if compat is not None:
        for setting in compat.findall(qn("w:compatSetting")):
            if setting.get(qn("w:name")) == "compatibilityMode":
                setting.set(qn("w:val"), "15")
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    if sources is None:
        _word_body(doc, markdown)
    else:
        doc.add_heading("Weekly Reading Collection", 0)
        doc.add_heading("Edition Summary", 1)
        _word_body(doc, summary, link_citations=True)
        doc.add_heading("Contents", 1)
        for number, source in enumerate(sources, 1):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.2)
            paragraph.add_run(f"{number:02}.  ").bold = True
            _word_link(paragraph, source_title(source), f"source-{number}")
        articles = doc.add_heading("Complete Articles", 1)
        articles.paragraph_format.page_break_before = True
        for number, source in enumerate(sources, 1):
            heading = doc.add_heading(f"Source {number}: {source_title(source)}", 2)
            if number > 1:
                heading.paragraph_format.page_break_before = True
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(number))
            start.set(qn("w:name"), f"source-{number}")
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(number))
            heading._p.insert(1, start)
            heading._p.append(end)
            _word_body(doc, source.content, skip_title=source_title(source))
        references = doc.add_heading("Original Sources", 1)
        references.paragraph_format.page_break_before = True
        for number, source in enumerate(sources, 1):
            paragraph = doc.add_paragraph()
            paragraph.add_run(f"{number:02}.  {source_title(source)}  -  ")
            if urlsplit(source.reference).scheme in {"http", "https"}:
                _word_link(paragraph, "Open original article", source.reference, external=True)
            else:
                paragraph.add_run(f"Local input: inputs/{source.label}")
    doc.save(path)


def _pdf_font() -> str:
    candidates = [
        Path(os.environ.get("WINDIR", "C:\\Windows")) / "Fonts" / "arial.ttf",
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.is_file():
            pdfmetrics.registerFont(TTFont("DocHarness", str(candidate)))
            return "DocHarness"
    return "Helvetica"


class LinkedPdf(SimpleDocTemplate):
    def afterFlowable(self, flowable) -> None:
        anchor = getattr(flowable, "_source_anchor", None)
        if anchor:
            self.canv.bookmarkPage(anchor)
            self.canv.addOutlineEntry(getattr(flowable, "_source_title"), anchor, level=0)


def write_pdf(
    markdown: str, path: Path, sources: list[IndexedSource] | None = None,
    summary: str = "",
) -> None:
    font = _pdf_font()
    base = ParagraphStyle(
        "Body", parent=getSampleStyleSheet()["BodyText"], fontName=font,
        fontSize=10, leading=15, textColor=colors.HexColor("#233047"), spaceAfter=9,
    )
    styles = {
        1: ParagraphStyle("Title", parent=base, fontSize=22, leading=27, spaceBefore=5, spaceAfter=20, textColor=colors.HexColor("#18264A")),
        2: ParagraphStyle("Heading", parent=base, fontSize=14, leading=19, spaceBefore=14, spaceAfter=9, textColor=colors.HexColor("#2155BC")),
        3: ParagraphStyle("Subheading", parent=base, fontSize=11, leading=16, spaceBefore=10, spaceAfter=6, textColor=colors.HexColor("#2155BC")),
    }
    def add_text(story: list, text: str, link_citations: bool = False) -> None:
        for kind, level, item in lines(text):
            item = plain(item)
            chunks = []
            position = 0
            for match in re.finditer(r"https?://[^\s<>]+", item):
                url = match.group().rstrip(".,)")
                chunks.append(html.escape(item[position:match.start()]))
                chunks.append(f'<link href="{html.escape(url, quote=True)}" color="#2155BC">{html.escape(url)}</link>')
                position = match.start() + len(url)
            chunks.append(html.escape(item[position:]))
            safe = "".join(chunks)
            if link_citations:
                safe = re.sub(
                    r"\[(\d+)\]",
                    lambda match: f'<link href="#source-{match.group(1)}" color="#2155BC">{match.group()}</link>',
                    safe,
                )
            story.append(Paragraph(
                safe if kind == "heading" else ("&bull; " if kind == "bullet" else "") + safe,
                styles[level] if kind == "heading" else base,
            ))

    story = []
    if sources is None:
        add_text(story, markdown)
    else:
        story.append(Paragraph("Weekly Reading Collection", styles[1]))
        story.append(Paragraph("Edition Summary", styles[2]))
        add_text(story, summary, link_citations=True)
        story.append(Paragraph("Contents", styles[2]))
        for number, source in enumerate(sources, 1):
            title = html.escape(f"Source {number}: {source_title(source)}")
            story.append(Paragraph(f'<link href="#source-{number}" color="#2155BC">{title}</link>', base))
        story.append(PageBreak())
        story.append(Paragraph("Complete Articles", styles[2]))
        for number, source in enumerate(sources, 1):
            if number > 1:
                story.append(PageBreak())
            title = f"Source {number}: {source_title(source)}"
            heading = Paragraph(html.escape(title), styles[2])
            heading._source_anchor = f"source-{number}"
            heading._source_title = title
            story.append(heading)
            add_text(story, source.content)
        story.append(PageBreak())
        story.append(Paragraph("Original Sources", styles[2]))
        for number, source in enumerate(sources, 1):
            title = html.escape(f"{number}. {source_title(source)}")
            if urlsplit(source.reference).scheme in {"http", "https"}:
                url = html.escape(source.reference, quote=True)
                story.append(Paragraph(f'<link href="{url}" color="#2155BC">{title}</link>', base))
                story.append(Paragraph(html.escape(source.reference), base))
            else:
                story.append(Paragraph(f"{title} — Local input: inputs/{html.escape(source.label)}", base))
    if not story:
        raise ValueError("The draft has no content.")
    doc = LinkedPdf(str(path), pagesize=(595, 842), rightMargin=55, leftMargin=55, topMargin=55, bottomMargin=55)
    doc.build(story)


def restore_approved(workspace: Workspace) -> bool:
    if workspace.state.get("pending"):
        return True
    if workspace.state.get("stage") != "Exported" or not workspace.state.get("last_output"):
        return False
    previous_output = Path(workspace.state["last_output"]).resolve()
    if not previous_output.is_relative_to(workspace.output) or not previous_output.is_file():
        raise ValueError("The previous output is missing. Run Build before exporting again.")
    draft_path = workspace.cache / "draft.md"
    if not draft_path.is_file():
        raise ValueError("The previous reviewed draft is missing. Run Build again.")
    draft = draft_path.read_text(encoding="utf-8")
    try:
        intro_with_links = draft.split("## Edition Summary\n\n", 1)[1].split("\n\n## Contents\n", 1)[0]
    except IndexError as error:
        raise ValueError("The previous draft cannot be reused. Run Build once.") from error
    intro = re.sub(
        r"\[(\d+)\]\(#source-\1\)", lambda match: f"[{match.group(1)}]", intro_with_links,
    )
    items = workspace.sources()
    sources: list[IndexedSource] = []
    snapshots = []
    for number, item in enumerate(items, 1):
        start = f'<a id="source-{number}"></a>\n### Source {number}:'
        if start not in draft:
            raise ValueError("The previous draft no longer matches its source list. Run Build again.")
        section = draft.split(start, 1)[1].split("\n\n", 1)[1]
        end = f'<a id="source-{number + 1}"></a>' if number < len(items) else "## Original Sources"
        if end not in section:
            raise ValueError("The previous draft is incomplete. Run Build again.")
        content = section.split(end, 1)[0].strip("\n")
        reference = item.get("url", item["label"])
        found = None
        for snapshot in workspace.cache.glob(f"source-{number:03d}-*.txt"):
            text = snapshot.read_text(encoding="utf-8")
            if text.strip("\n") == content:
                found = snapshot
                break
        if found is None:
            raise ValueError("A previous source snapshot is missing. Run Build again.")
        if item["kind"] == "file" and read_file(item["path"]).strip("\n") != content:
            raise ValueError("An input file changed since approval. Run Build again.")
        text = found.read_text(encoding="utf-8")
        digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
        sources.append(IndexedSource(item["label"], reference, text))
        snapshots.append({
            "label": item["label"], "reference": reference,
            "snapshot": found.name, "sha256": digest,
        })
    quality = review(intro, sources, workspace.state["settings"]["target_words"])
    if quality.issues:
        raise ValueError("The previous summary no longer passes citation review. Run Build again.")
    normalized = compose(intro, sources)
    if normalized != draft:
        backup = workspace.cache / f"draft-before-reuse-{hashlib.sha256(draft.encode()).hexdigest()[:12]}.md"
        if not backup.exists():
            atomic_text(backup, draft)
        atomic_text(draft_path, normalized)
    workspace.state["pending"] = {
        "fingerprint": workspace.fingerprint(),
        "sources": snapshots,
        "intro": intro,
        "draft_sha256": hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
        "review": {
            "issues": quality.issues, "warnings": quality.warnings, "word_count": quality.word_count,
        },
        "approved": True,
    }
    workspace.state["last_format"] = previous_output.suffix.lstrip(".").replace("md", "markdown")
    workspace.save()
    return True


def export(workspace: Workspace) -> str:
    restore_approved(workspace)
    pending = workspace.state.get("pending")
    if not pending:
        raise ValueError("No draft is waiting for approval. Run /build first.")
    if pending["fingerprint"] != workspace.fingerprint():
        workspace.invalidate()
        workspace.save()
        raise ValueError("The sources or settings changed since the draft. Run /build again.")
    if pending["review"]["issues"]:
        raise ValueError("Quality gate blocked export: " + "; ".join(pending["review"]["issues"]) + ". Use /revise.")
    draft = (workspace.cache / "draft.md").read_text(encoding="utf-8")
    if hashlib.sha256(draft.encode("utf-8")).hexdigest() != pending.get("draft_sha256"):
        raise ValueError("The reviewed draft changed. Run /build or /revise before exporting.")
    if "sources" not in pending:
        raise ValueError("This draft predates full-text mode. Run /build again.")
    sources = load_snapshots(workspace, pending["sources"])
    if draft != compose(pending["intro"], sources):
        raise ValueError("The draft no longer matches the complete source texts. Run /build again.")
    fmt = workspace.state["settings"]["format"]
    extension = {"markdown": ".md", "pdf": ".pdf", "docx": ".docx"}[fmt]
    stem = f"reading-list-{date.today().isoformat()}"
    path = workspace.output / f"{stem}{extension}"
    number = 2
    while path.exists():
        path = workspace.output / f"{stem}-{number}{extension}"
        number += 1

    def write(temporary: Path) -> None:
        if fmt == "markdown":
            temporary.write_text(draft, encoding="utf-8")
        elif fmt == "pdf":
            write_pdf(draft, temporary, sources, pending["intro"])
        else:
            write_docx(draft, temporary, sources, pending["intro"])
        max_kb = workspace.state["settings"]["max_file_kb"]
        if max_kb is not None and temporary.stat().st_size > max_kb * 1024:
            raise ValueError(f"Output exceeds the configured {max_kb} KB limit.")

    _temp_output(path, write)
    result = str(path)
    workspace.state["stage"] = "Exported"
    workspace.state["last_output"] = result
    workspace.state["last_format"] = fmt
    pending["approved"] = True
    workspace.save()
    return result
